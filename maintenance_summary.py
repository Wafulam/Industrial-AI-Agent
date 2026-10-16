
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

from machine_runtime import load_state
from maintenance_metrics import load_events

EXCEL_FILE = Path(__file__).parent / "motor_maintenance.xlsx"


def update_maintenance_summary():
    # Load machine runtime
    state = load_state()
    running_seconds = state["total_running_seconds"]

    if state["status"] == "RUNNING":
        started = datetime.fromisoformat(state["last_started"])
        running_seconds += (datetime.now() - started).total_seconds()

    running_hours = running_seconds / 3600

    # Load maintenance events
    events = load_events()

    failures = [
        event for event in events
        if event["event"] == "failure"
    ]

    repairs = [
        event for event in events
        if event["event"] == "repair"
    ]

    repair_durations = [
        event["repair_hours"] for event in repairs
    ]

    # Calculate MTTR
    mttr_hours = (
        sum(repair_durations) / len(repair_durations)
        if repair_durations else 0
    )

    # Calculate elapsed-time MTBF as a temporary simulation metric
    failure_times = sorted(
        datetime.fromisoformat(event["time"])
        for event in failures
    )

    intervals = [
        (failure_times[i] - failure_times[i - 1]).total_seconds() / 3600
        for i in range(1, len(failure_times))
    ]

    mtbf_hours = sum(intervals) / len(intervals) if intervals else 0

    # Simplified availability against an 8-hour shift
    scheduled_hours = 8
    availability = min(running_hours / scheduled_hours * 100, 100)

    # Open or create workbook
    if EXCEL_FILE.exists():
        workbook = load_workbook(EXCEL_FILE)
    else:
        from openpyxl import Workbook
        workbook = Workbook()

    # Create or refresh the summary worksheet
    if "Maintenance Summary" in workbook.sheetnames:
        sheet = workbook["Maintenance Summary"]
        sheet.delete_rows(1, sheet.max_row)
    else:
        sheet = workbook.create_sheet("Maintenance Summary")

    sheet.append(["INDUSTRIAL MAINTENANCE SUMMARY", "Value"])
    sheet.append(["Last Updated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    sheet.append(["Machine Status", state["status"]])
    sheet.append(["Accumulated Running Hours", round(running_hours, 4)])
    sheet.append(["Scheduled Shift Hours (simulation)", scheduled_hours])
    sheet.append(["Availability (%) - simulation", round(availability, 2)])
    sheet.append(["Confirmed Breakdowns", len(failures)])
    sheet.append(["Completed Repairs", len(repairs)])
    sheet.append(["MTTR (hours)", round(mttr_hours, 4) if repairs else "N/A"])
    sheet.append(["MTTR (minutes)", round(mttr_hours * 60, 2) if repairs else "N/A"])
    sheet.append(["MTBF (elapsed hours, simulation)", round(mtbf_hours, 4) if intervals else "N/A"])

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")

    sheet.column_dimensions["A"].width = 42
    sheet.column_dimensions["B"].width = 30
    sheet.freeze_panes = "A2"

    workbook.save(EXCEL_FILE)

    print("Maintenance Summary updated in motor_maintenance.xlsx")
