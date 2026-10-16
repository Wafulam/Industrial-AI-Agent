
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

EXCEL_FILE = Path(__file__).parent / "motor_maintenance.xlsx"


def format_sheet(sheet):
    """Format worksheet headers and columns."""

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )
        cell.alignment = Alignment(horizontal="center")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for column in sheet.columns:
        column_letter = column[0].column_letter
        max_length = max(
            len(str(cell.value or "")) for cell in column
        )
        sheet.column_dimensions[column_letter].width = min(
            max_length + 3, 45
        )


def get_workbook():
    """Open the existing workbook or create a new one."""

    if EXCEL_FILE.exists():
        workbook = load_workbook(EXCEL_FILE)
    else:
        workbook = Workbook()
        workbook.active.title = "Motor History"

    if "Motor History" not in workbook.sheetnames:
        sheet = workbook.create_sheet("Motor History")
        sheet.append([
            "Timestamp",
            "Temperature (C)",
            "Current (A)",
            "Vibration",
            "Speed (RPM)",
            "AI Condition",
            "AI Confidence (%)",
            "Recommendation"
        ])
        format_sheet(sheet)

    if "Maintenance Events" not in workbook.sheetnames:
        sheet = workbook.create_sheet("Maintenance Events")
        sheet.append([
            "Timestamp",
            "Event",
            "Details",
            "Duration (hours)"
        ])
        format_sheet(sheet)

    return workbook


def save_to_excel(motor_data, condition, confidence, recommendation):
    """Save motor readings and AI results."""

    workbook = get_workbook()
    sheet = workbook["Motor History"]

    if sheet.max_row == 1 and sheet["A1"].value is None:
        sheet.append([
            "Timestamp",
            "Temperature (C)",
            "Current (A)",
            "Vibration",
            "Speed (RPM)",
            "AI Condition",
            "AI Confidence (%)",
            "Recommendation"
        ])

    predicted_confidence = confidence.get(condition, 0)

    sheet.append([
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        motor_data["temperature"],
        motor_data["current"],
        motor_data["vibration"],
        motor_data["speed"],
        condition,
        predicted_confidence,
        recommendation
    ])


    format_sheet(sheet)
    workbook.save(EXCEL_FILE)
    print("Motor readings saved to Excel.")


def save_maintenance_event(event, details="", duration_hours=None):
    """Save a breakdown or repair event."""

    workbook = get_workbook()
    sheet = workbook["Maintenance Events"]

    sheet.append([
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        event,
        details,
        duration_hours
    ])

    format_sheet(sheet)
    workbook.save(EXCEL_FILE)
    print(f"Maintenance event saved to Excel: {event}")
    
def save_alert_to_excel(condition, message):
    """Save a motor alert to the Maintenance Events worksheet."""

    workbook = get_workbook()
    sheet = workbook["Maintenance Events"]

    sheet.append([
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        f"AI {condition} Alert",
        message,
        None
    ])

    format_sheet(sheet)
    workbook.save(EXCEL_FILE)

    print(f"AI alert saved to Excel: {condition}")

    
def show_alert_history():
    """Display previously logged AI alerts."""

    if not EXCEL_FILE.exists():
        print("No Excel maintenance records found.")
        return

    workbook = load_workbook(EXCEL_FILE)
    
    if "Maintenance Events" not in workbook.sheetnames:
        print("No maintenance events worksheet found.")
        return

    sheet = workbook["Maintenance Events"]

    print("\n--- AI ALERT HISTORY ---")
    found = False

    for row in sheet.iter_rows(min_row=2, values_only=True):
        timestamp, event, details, duration = row[:4]

        if event in ("AI Warning Alert", "AI Fault Alert"):
            print(f"Time: {timestamp}")
            print(f"Event: {event}")
            print(f"Details: {details}")
            print("-" * 35)
            found = True

    if not found:
        print("No AI alerts have been logged yet.")


